"""
Geradores de relatorio em PDF (reportlab) e DOCX (python-docx).
Recebem o mesmo dicionario de contexto produzido pelo conftest.
"""
from pathlib import Path
import datetime as dt

# ===================== PDF =====================

def write_pdf(ctx: dict, out_path: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
    )

    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Heading1"], textColor=colors.HexColor("#6C63FF"))
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], textColor=colors.HexColor("#333366"))
    normal = styles["BodyText"]
    small = ParagraphStyle("small", parent=normal, fontSize=8, leading=10)

    doc = SimpleDocTemplate(
        str(out_path), pagesize=A4,
        leftMargin=2 * cm, rightMargin=2 * cm,
        topMargin=1.8 * cm, bottomMargin=1.8 * cm,
        title="Relatorio de Testes - tiResolve",
    )

    story = []
    story.append(Paragraph("Relatorio de Testes Web -> Backend", h1))
    story.append(Paragraph("tiResolve - Centro de Informatica", normal))
    story.append(Spacer(1, 0.4 * cm))

    meta = [
        ["Data",       ctx["data"]],
        ["URL alvo",   ctx["base_url"]],
        ["Run ID",     ctx["run_id"]],
        ["Modo",       ctx["modo"]],
    ]
    t = Table(meta, colWidths=[3.5 * cm, 12 * cm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EEEEFF")),
        ("FONTNAME",   (0, 0), (-1, -1), "Helvetica"),
        ("FONTSIZE",   (0, 0), (-1, -1), 9),
        ("BOX",        (0, 0), (-1, -1), 0.4, colors.grey),
        ("INNERGRID",  (0, 0), (-1, -1), 0.2, colors.grey),
        ("VALIGN",     (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story.append(t)
    story.append(Spacer(1, 0.5 * cm))

    # Resumo
    story.append(Paragraph("Resumo", h2))
    s = ctx["summary"]
    resumo = [
        ["Total", "Passou", "Falhou", "Skipped", "Duracao"],
        [s["total"], s["passed"], s["failed"], s["skipped"], f"{s['duration_s']:.2f}s"],
    ]
    t = Table(resumo, colWidths=[3 * cm] * 5)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#6C63FF")),
        ("TEXTCOLOR",  (0, 0), (-1, 0), colors.white),
        ("FONTNAME",   (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN",      (0, 0), (-1, -1), "CENTER"),
        ("FONTSIZE",   (0, 0), (-1, -1), 10),
        ("BOX",        (0, 0), (-1, -1), 0.4, colors.grey),
        ("GRID",       (0, 0), (-1, -1), 0.2, colors.grey),
        ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#F4F4FF")),
    ]))
    story.append(t)
    story.append(Spacer(1, 0.5 * cm))

    # Testes
    story.append(Paragraph("Resultado por caso de teste", h2))
    data = [["#", "Teste", "Status", "Tempo (s)"]]
    for i, r in enumerate(ctx["tests"], 1):
        nice = r["nodeid"].split("::")[-1]
        data.append([str(i), nice, r["outcome"], f"{r['duration']:.2f}"])
    t = Table(data, colWidths=[1 * cm, 11 * cm, 2.5 * cm, 2.5 * cm], repeatRows=1)
    cell_styles = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#333366")),
        ("TEXTCOLOR",  (0, 0), (-1, 0), colors.white),
        ("FONTNAME",   (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE",   (0, 0), (-1, -1), 9),
        ("ALIGN",      (0, 1), (0, -1), "CENTER"),
        ("ALIGN",      (2, 1), (-1, -1), "CENTER"),
        ("GRID",       (0, 0), (-1, -1), 0.2, colors.grey),
    ]
    for i, r in enumerate(ctx["tests"], 1):
        if r["outcome"] == "PASSED":
            cell_styles.append(("TEXTCOLOR", (2, i), (2, i), colors.HexColor("#1A8917")))
        elif r["outcome"] == "FAILED":
            cell_styles.append(("TEXTCOLOR", (2, i), (2, i), colors.HexColor("#C62828")))
            cell_styles.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#FFEBEE")))
    t.setStyle(TableStyle(cell_styles))
    story.append(t)
    story.append(Spacer(1, 0.5 * cm))

    # Endpoints
    story.append(Paragraph("Endpoints do backend acionados", h2))
    if ctx["endpoints"]:
        data = [["Endpoint", "Total", "2xx/3xx", "Erros"]]
        for key in sorted(ctx["endpoints"].keys()):
            v = ctx["endpoints"][key]
            data.append([key, str(v["total"]), str(v["ok"]), str(v["fail"])])
        t = Table(data, colWidths=[10 * cm, 2 * cm, 2.5 * cm, 2.5 * cm], repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#333366")),
            ("TEXTCOLOR",  (0, 0), (-1, 0), colors.white),
            ("FONTNAME",   (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE",   (0, 0), (-1, -1), 8),
            ("ALIGN",      (1, 1), (-1, -1), "CENTER"),
            ("GRID",       (0, 0), (-1, -1), 0.2, colors.grey),
        ]))
        story.append(t)
    else:
        story.append(Paragraph("(sem logs do backend disponiveis)", normal))
    story.append(Spacer(1, 0.5 * cm))

    # Cobertura
    story.append(Paragraph("Cobertura por modulo do backend", h2))
    cob = [["Modulo", "Endpoints", "Status"]]
    for nome, lst in ctx["modulos"].items():
        cob.append([nome, str(len(lst)), "OK" if lst else "nao exercitado"])
    t = Table(cob, colWidths=[8 * cm, 3 * cm, 6 * cm], repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#333366")),
        ("TEXTCOLOR",  (0, 0), (-1, 0), colors.white),
        ("FONTNAME",   (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE",   (0, 0), (-1, -1), 9),
        ("ALIGN",      (1, 1), (1, -1), "CENTER"),
        ("GRID",       (0, 0), (-1, -1), 0.2, colors.grey),
    ]))
    story.append(t)

    # Detalhamento "o que / como / endpoints" por teste
    story.append(PageBreak())
    story.append(Paragraph("Detalhamento dos testes (o que e como)", h2))
    casos = ctx.get("casos") or {}
    bullet = ParagraphStyle("bullet", parent=normal, leftIndent=14, bulletIndent=4, spaceAfter=2)
    for i, r in enumerate(ctx["tests"], 1):
        nice = r["nodeid"].split("::")[-1]
        meta = casos.get(nice, {})
        story.append(Paragraph(
            f"<b>{i}. {nice}</b> &mdash; <font color='{'#1A8917' if r['outcome']=='PASSED' else '#C62828'}'>{r['outcome']}</font> ({r['duration']:.2f}s)",
            normal,
        ))
        if meta:
            story.append(Paragraph(f"<b>O que e testado:</b> {meta.get('descricao','-')}", bullet))
            story.append(Paragraph(f"<b>Como (passos na WEB):</b> {meta.get('como','-')}", bullet))
            eps = meta.get("endpoints") or []
            if eps:
                story.append(Paragraph("<b>Endpoints do backend exercitados:</b>", bullet))
                for ep in eps:
                    story.append(Paragraph(f"&bull; <font face='Courier'>{ep}</font>", bullet))
            vals = meta.get("validacoes") or []
            if vals:
                story.append(Paragraph("<b>Validacoes (asserts):</b>", bullet))
                for v in vals:
                    story.append(Paragraph(f"&bull; {v}", bullet))
        else:
            story.append(Paragraph("<i>(sem metadados no catalogo CASOS)</i>", bullet))
        story.append(Spacer(1, 0.25 * cm))

    if any(r["outcome"] == "FAILED" for r in ctx["tests"]):
        story.append(PageBreak())
        story.append(Paragraph("Falhas (detalhes)", h2))
        for r in ctx["tests"]:
            if r["outcome"] == "FAILED":
                story.append(Paragraph(f"<b>{r['nodeid']}</b>", normal))
                trunc = (r["longrepr"] or "")[:1500].replace("<", "&lt;").replace(">", "&gt;")
                story.append(Paragraph(f"<font face='Courier' size='8'>{trunc}</font>", small))
                story.append(Spacer(1, 0.3 * cm))

    doc.build(story)


def write_docx(ctx: dict, out_path: Path) -> None:
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.table import WD_ALIGN_VERTICAL

    doc = Document()
    # Margens
    for section in doc.sections:
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)

    title = doc.add_heading("Relatorio de Testes Web -> Backend", level=0)
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x6C, 0x63, 0xFF)
    doc.add_paragraph("tiResolve - Centro de Informatica")

    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.style = "Light Grid Accent 1"
    pares = [("Data", ctx["data"]), ("URL alvo", ctx["base_url"]),
             ("Run ID", ctx["run_id"]), ("Modo", ctx["modo"])]
    for i, (k, v) in enumerate(pares):
        meta_table.rows[i].cells[0].text = k
        meta_table.rows[i].cells[1].text = str(v)

    doc.add_heading("Resumo", level=1)
    s = ctx["summary"]
    rt = doc.add_table(rows=2, cols=5)
    rt.style = "Light Grid Accent 1"
    hdrs = ["Total", "Passou", "Falhou", "Skipped", "Duracao"]
    vals = [s["total"], s["passed"], s["failed"], s["skipped"], f"{s['duration_s']:.2f}s"]
    for i, h in enumerate(hdrs):
        rt.rows[0].cells[i].text = h
        rt.rows[1].cells[i].text = str(vals[i])

    doc.add_heading("Resultado por caso de teste", level=1)
    tt = doc.add_table(rows=1, cols=4)
    tt.style = "Light Grid Accent 1"
    for i, h in enumerate(["#", "Teste", "Status", "Tempo (s)"]):
        tt.rows[0].cells[i].text = h
    for i, r in enumerate(ctx["tests"], 1):
        row = tt.add_row().cells
        row[0].text = str(i)
        row[1].text = r["nodeid"].split("::")[-1]
        row[2].text = r["outcome"]
        row[3].text = f"{r['duration']:.2f}"
        cor = (0x1A, 0x89, 0x17) if r["outcome"] == "PASSED" else \
              (0xC6, 0x28, 0x28) if r["outcome"] == "FAILED" else (0x99, 0x99, 0x99)
        for p in row[2].paragraphs:
            for run in p.runs:
                run.font.color.rgb = RGBColor(*cor)
                run.font.bold = True

    doc.add_heading("Endpoints do backend acionados", level=1)
    if ctx["endpoints"]:
        et = doc.add_table(rows=1, cols=4)
        et.style = "Light Grid Accent 1"
        for i, h in enumerate(["Endpoint", "Total", "2xx/3xx", "Erros"]):
            et.rows[0].cells[i].text = h
        for key in sorted(ctx["endpoints"].keys()):
            v = ctx["endpoints"][key]
            row = et.add_row().cells
            row[0].text = key
            row[1].text = str(v["total"])
            row[2].text = str(v["ok"])
            row[3].text = str(v["fail"])
    else:
        doc.add_paragraph("(sem logs do backend disponiveis)")

    doc.add_heading("Cobertura por modulo do backend", level=1)
    ct = doc.add_table(rows=1, cols=3)
    ct.style = "Light Grid Accent 1"
    for i, h in enumerate(["Modulo", "Endpoints", "Status"]):
        ct.rows[0].cells[i].text = h
    for nome, lst in ctx["modulos"].items():
        row = ct.add_row().cells
        row[0].text = nome
        row[1].text = str(len(lst))
        row[2].text = "OK" if lst else "nao exercitado"

    # Detalhamento "o que / como / endpoints" por teste
    doc.add_page_break()
    doc.add_heading("Detalhamento dos testes (o que e como)", level=1)
    casos = ctx.get("casos") or {}
    for i, r in enumerate(ctx["tests"], 1):
        nice = r["nodeid"].split("::")[-1]
        meta = casos.get(nice, {})
        p = doc.add_paragraph()
        p.add_run(f"{i}. {nice} — ").bold = True
        st_run = p.add_run(r["outcome"])
        st_run.bold = True
        st_run.font.color.rgb = RGBColor(
            *((0x1A, 0x89, 0x17) if r["outcome"] == "PASSED" else
              (0xC6, 0x28, 0x28) if r["outcome"] == "FAILED" else (0x99, 0x99, 0x99))
        )
        p.add_run(f"  ({r['duration']:.2f}s)")

        if not meta:
            doc.add_paragraph("(sem metadados no catalogo CASOS)").italic = True
            continue
        p = doc.add_paragraph()
        p.add_run("O que e testado: ").bold = True
        p.add_run(meta.get("descricao", "-"))
        p = doc.add_paragraph()
        p.add_run("Como (passos na WEB): ").bold = True
        p.add_run(meta.get("como", "-"))
        eps = meta.get("endpoints") or []
        if eps:
            doc.add_paragraph("Endpoints do backend exercitados:").runs[0].bold = True
            for ep in eps:
                doc.add_paragraph(ep, style="List Bullet")
        vals = meta.get("validacoes") or []
        if vals:
            doc.add_paragraph("Validacoes (asserts):").runs[0].bold = True
            for v in vals:
                doc.add_paragraph(v, style="List Bullet")

    if any(r["outcome"] == "FAILED" for r in ctx["tests"]):
        doc.add_page_break()
        doc.add_heading("Falhas (detalhes)", level=1)
        for r in ctx["tests"]:
            if r["outcome"] == "FAILED":
                doc.add_paragraph(r["nodeid"]).runs[0].bold = True
                p = doc.add_paragraph()
                run = p.add_run((r["longrepr"] or "")[:2000])
                run.font.name = "Courier New"
                run.font.size = Pt(8)

    doc.save(str(out_path))
